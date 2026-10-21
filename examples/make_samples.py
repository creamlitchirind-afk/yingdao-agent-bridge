"""Generate fictitious, reproducible inputs; never package personal workbooks."""
from pathlib import Path
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

def main():
    target=Path(__file__).resolve().parents[1]/'.local'/'samples'
    target.mkdir(parents=True,exist_ok=True)
    book=Workbook(); sheet=book.active
    sheet.append(['联系方式','确认状态','姓名','报名编号','报名项目'])
    for row in [
        ['00010001','已确认','林琪','A001','绘画'],
        ['', '已确认','陈晨','A002','书法'],
        ['00010001','已确认','林琪','A001','绘画'],
        ['00010003','已取消','沈川','A003','书法'],
        ['00010004','已确认','周可','A004','摄影']]: sheet.append(row)
    book.save(target/'registrations.xlsx')
    book=Workbook(); sheet=book.active
    sheet['A1']='虚构活动名单'; sheet['A2']='样例，不含真实个人信息'
    for i,value in enumerate(['项目','联系电话','姓名','编号','备注'],1): sheet.cell(3,i,value)
    for row in range(4,14): sheet.cell(row,5,'人工备注保留')
    sheet['A15']='审核人：待填写'
    for column in 'ABCDE': sheet.column_dimensions[column].width=22
    book.save(target/'template.xlsx')
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    pdf=canvas.Canvas(str(target/'orders.pdf'))
    for order in [('DEMO-001','示例文具供应商','2026-10-01','120.50'),('DEMO-002','示例纸品供应商','2026-10-02','80.00')]:
        pdf.setFont('STSong-Light',14)
        for index,(label,value) in enumerate(zip(['订单编号','供应商','日期','金额'],order)):
            pdf.drawString(60,760-index*35,label+'：'+value)
        pdf.showPage()
    pdf.save()
    print(target)

if __name__=='__main__': main()
