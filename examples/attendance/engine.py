"""Local Excel processor invoked by Yingdao. Python 3.11 + openpyxl.
Only designated submitted folders are read. Fail closed on ambiguous inputs.
"""
import argparse,copy,hashlib,json,re,os,unicodedata,math
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font,Alignment,Border,Side
from openpyxl.utils import column_index_from_string

def lesson_events(s,row,cols,groups=None):
    """Combine late/early only inside an explicitly supplied lesson boundary."""
    spans=groups if groups is not None else [[c] for c in cols]
    for span in spans:
        late_cells=[]
        for c in span:
            cell=s.cell(row,c);mark=normalized(cell.value).replace(' ','')
            if not mark or mark in ('△','R','r'):continue
            if mark in ('×','X','x'):yield 1,cell.coordinate
            elif all(ch in '○Oo⚪#' for ch in mark):late_cells.append(cell.coordinate)
            else:raise ValueError(f'{cell.coordinate}:未知考勤符号 {mark}')
        if late_cells:yield .5,','.join(late_cells)

def normalized(v):
    return unicodedata.normalize('NFKC',str(v or '')).replace('_x000D_','').strip()

def locate(s,label):
    hits=[]
    for row in s.iter_rows(min_row=1,max_row=min(s.max_row,6)):
        for c in row:
            if normalized(c.value).replace(' ','').rstrip(':：')==label:
                # Accommodate multi-column merged field labels.
                end=next((r.max_col for r in s.merged_cells.ranges if c.coordinate in r),c.column)
                for cc in range(end+1,min(s.max_column,end+4)+1):
                    v=s.cell(c.row,cc).value
                    if v is not None:hits.append(v);break
    return hits

def identify(f,s,colleges,week,kind):
    nf=[c for c in colleges if c in f.stem]
    inside=[]
    if kind=='attendance':inside=[normalized(v) for v in locate(s,'学院') if normalized(v) in colleges]
    else:
        headers=[c for row in s.iter_rows(min_row=1,max_row=min(s.max_row,5)) for c in row if normalized(c.value)=='学院']
        for h in headers:
            for r in range(h.row+1,s.max_row+1):
                v=normalized(s.cell(r,h.column).value)
                if v in colleges:inside.append(v)
    both=set(nf+inside)
    if len(both)!=1:raise ValueError(f'{f.name}: 学院缺失或冲突 {sorted(both)}')
    # Detect numeric demo week everywhere in header and file, ignore lesson numbers.
    header=' '.join(str(c.value or '') for row in s.iter_rows(min_row=1,max_row=min(s.max_row,3)) for c in row)
    weeks={int(x) for x in re.findall(r'第\s*(\d+)\s*周',f.stem+' '+header)}
    if weeks!={week}:raise ValueError(f'{f.name}: 周次不明确或不一致 {weeks}, 目标{week}')
    return both.pop(),not bool(nf)

def parse_week(root,colleges,week,progress=None):
    notify=progress or (lambda message:None)
    events=[];submitted=set();files=[];fallback=[];identities=set();errors=[]
    received={'attendance':set(),'late':set()}
    schedule_path=root/'课时分组.json'
    lesson_groups=None
    if schedule_path.exists():
        lesson_groups=[[column_index_from_string(c) for c in group] for group in json.loads(schedule_path.read_text(encoding='utf-8'))['groups']]
    for folder,kind in [('考勤表','attendance'),('晚点名','late')]:
        directory=root/'02_已提交材料'/f'第{week}周'/folder
        if not directory.is_dir():raise ValueError('输入目录不存在: '+str(directory))
        for f in sorted(directory.glob('*.xlsx')):
            if f.name.startswith('~$'):continue
            try:
                w=load_workbook(f,data_only=True);s=w.active
                college,fb=identify(f,s,colleges,week,kind)
                if fb:fallback.append(f.name)
                digest=hashlib.sha256(f.read_bytes()).hexdigest()
                if any(x['sha256']==digest for x in files):raise ValueError('同周出现完全重复文件: '+f.name)
                files.append({'file':str(f),'sha256':digest});submitted.add(college);received[kind].add(college)
                def emit(klass,name,hours,cell,event):
                    key=(college,klass,name);identities.add(key)
                    events.append({'college':college,'class':klass,'name':name,'hours':hours,'kind':event,'file':str(f),'sheet':s.title,'cell':cell,'week':week})
                if kind=='attendance':
                    # Locate name/class header, lesson columns bounded by actual lesson-number row.
                    hr=next((r for r in range(1,8) if normalized(s.cell(r,1).value)=='班级' and normalized(s.cell(r,2).value)=='姓名'),None)
                    if hr is None:raise ValueError('未找到班级/姓名表头')
                    periodrow=next((r for r in range(hr+1,hr+4) if sum(isinstance(s.cell(r,c).value,(int,float)) for c in range(3,s.max_column+1))>5),None)
                    if periodrow is None:raise ValueError('未找到课时表头')
                    cols=[c for c in range(3,s.max_column+1) if isinstance(s.cell(periodrow,c).value,(int,float))]
                    if lesson_groups is not None:
                        flat=[c for group in lesson_groups for c in group]
                        if len(flat)!=len(set(flat)) or set(flat)!=set(cols):raise ValueError('课时分组与表头不一致')
                    for r in range(periodrow+1,s.max_row+1):
                        klass=normalized(s.cell(r,1).value);name=normalized(s.cell(r,2).value)
                        if not name:continue
                        if not klass:raise ValueError(f'第{r}行班级为空')
                        identities.add((college,klass,name))
                        for h,cell in lesson_events(s,r,cols,lesson_groups):
                            emit(klass,name,h,cell,'课堂')
                else:
                    # Original layout: G count, H combined names; class must be explicit.
                    for r in range(3,s.max_row+1):
                        if normalized(s.cell(r,2).value)!=college:continue
                        count=s.cell(r,7).value;actual=s.cell(r,4).value;expected=s.cell(r,3).value;leave=s.cell(r,5).value or 0
                        if not isinstance(count,(int,float)) or count<0 or int(count)!=count:raise ValueError('未签到人数无效')
                        if not all(isinstance(x,(int,float)) for x in [actual,expected,leave]) or actual+count+leave!=expected:raise ValueError('签到人数与应到人数不符')
                        raw=normalized(s.cell(r,8).value)
                        people=[] if raw in ('','无','0') else [normalized(x) for x in raw.splitlines() if normalized(x)]
                        if len(people)!=count:raise ValueError('H列名单人数与G列不一致')
                        for line in people:
                            bits=line.split()
                            if len(bits)!=2:raise ValueError('名单须明确班级和姓名，无法判定: '+line)
                            emit(bits[0],bits[1],3,f'H{r}','晚点名')
                w.close()
                notify(f'[材料读取完成] {college}第{week}周{folder}：{f.name}')
            except Exception as exc:
                errors.append(f'{f.name}: {exc}')
                notify(f'[材料错误] 第{week}周 {f.name}：{exc}')
    if errors:raise ValueError('\n'.join(errors))
    return {'week':week,'submitted':sorted(submitted),'received':{k:sorted(v) for k,v in received.items()},'missing_late':sorted(received['attendance']-received['late']),'events':events,'files':files,'filename_fallback':fallback,'identities':[list(k) for k in sorted(identities)]}

def aggregate(data):
    groups={};submitted=set()
    for d in data:
        submitted.update(d['submitted'])
        for e in d['events']:
            key=(e['college'],e['class'],e['name'])
            v=groups.setdefault(key,{'hours':0,'late':0,'evidence':[]})
            v['hours']+=e['hours'];v['late']+=int(e['kind']=='晚点名');v['evidence'].append(e)
    return groups,submitted

def copy_style(source,target):
    # Style IDs are workbook-local; copying _style across files corrupts fonts.
    for attr in ('font','fill','border','alignment','protection'):
        setattr(target,attr,copy.copy(getattr(source,attr)))
    target.number_format=source.number_format

def fill(s,reference,groups,submitted,colleges,title,cumulative,missing_late=None,complete=None):
    missing_late=set(missing_late or [])
    complete=set(complete or [])
    end=6 if cumulative else 5
    for m in list(s.merged_cells.ranges):s.unmerge_cells(str(m))
    if s.max_row>2:s.delete_rows(3,s.max_row-2)
    for key in list(s.row_dimensions):
        if key>2:del s.row_dimensions[key]
    s['A1']=title;s.merge_cells(start_row=1,start_column=1,end_row=1,end_column=end)
    thin=Side(style='thin',color='000000');grid=Border(left=thin,right=thin,top=thin,bottom=thin)
    ordered=sorted(colleges,key=lambda c:(1 if any(k[0]==c for k in groups) else 0 if c in complete else 2,sum(k[0]==c for k in groups),colleges.index(c)))
    row=3
    for college in ordered:
        records=sorted([(k,v) for k,v in groups.items() if k[0]==college],key=lambda kv:(-kv[1]['hours'],kv[0][2],kv[0][1]))
        start=row
        if not records:records=[None]
        for record in records:
            for col in range(1,end+1):
                cell=s.cell(row,col);copy_style(reference.cell(4,col),cell);cell.border=grid
                cell.font=Font(name='宋体',size=11 if col==4 else 12,color='000000');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            s.row_dimensions[row].height=reference.row_dimensions[4].height or reference.sheet_format.defaultRowHeight
            s.cell(row,end).font=Font(name='宋体',size=12,color='FF0000')
            if record:
                k,v=record;s.cell(row,2,k[1]);s.cell(row,3,k[2]);s.cell(row,4,v['hours']);s.cell(row,4).number_format='General'
                notes=[]
                if v['late']:notes.append(f"晚点名未签到{v['late']}次")
                if college in missing_late:notes.append('未报送，该名单为学工处考勤结果')
                if notes:s.cell(row,end,'；'.join(notes))
            else:
                s.cell(row,2,'全勤' if college in complete else '未及时报送');s.cell(row,2).font=Font(name='宋体',size=12,color='000000' if college in complete else 'FF0000')
                s.merge_cells(start_row=row,start_column=2,end_row=row,end_column=4 if college in complete else end)
            # Preserve template column widths while allowing long classes/notes to wrap visibly.
            for col in (2,3,end):
                cell=s.cell(row,col)
                if not cell.value:continue
                width=s.column_dimensions[cell.column_letter].width or 10
                capacity=max(1,width*7-8)
                size=cell.font.sz or 12
                lines=sum(max(1,math.ceil(sum(size*4/3*(1 if unicodedata.east_asian_width(ch) in 'WF' else .55) for ch in line)/capacity)) for line in str(cell.value).split('\n'))
                if lines>1:s.row_dimensions[row].height=max(s.row_dimensions[row].height or 18,lines*size*1.25+3)
            row+=1
        s.cell(start,1,college)
        if row-start>1:s.merge_cells(start_row=start,start_column=1,end_row=row-1,end_column=1)
    s.cell(row,1,'注：迟到或早退两次算旷课一课时。');s.cell(row,1).font=Font(name='宋体',size=12,bold=True,color='C00000');s.cell(row,1).alignment=Alignment(horizontal='center',vertical='center');s.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end)
    s.print_area=f'A1:{chr(64+end)}{row}'
    return ordered

def run(root,output,week,progress=None):
    root=Path(root).resolve();output=Path(output).resolve()
    if output==root or root in output.parents:raise ValueError('结果目录必须与输入材料目录分开')
    output.mkdir(parents=True,exist_ok=True)
    blank=root/'01_起始总表/演示学期学生课堂纪律总表_空白.xlsx'
    weekly=root/'01_起始总表/周汇总空白模板.xlsx'
    base=load_workbook(blank);base.active=0;colleges=[base.active.cell(r,1).value for r in range(3,16)]
    if len(set(colleges))!=13:raise ValueError('初始总表学院名单不完整')
    statepath=output/'汇总台账.json';state=json.loads(statepath.read_text(encoding='utf-8')) if statepath.exists() else {'weeks':{}}
    data=parse_week(root,colleges,week,progress=progress);state['weeks'][str(week)]=data
    weeks=sorted(map(int,state['weeks']));groups,submitted=aggregate([state['weeks'][str(n)] for n in weeks])
    scope=f'第{weeks[0]}周' if len(weeks)==1 else f'第{weeks[0]}-{weeks[-1]}周'
    if weeks!=list(range(weeks[0],weeks[-1]+1)):raise ValueError('累计周次不连续，不能使用连续范围标题')
    ref=load_workbook(blank).worksheets[0]
    for s in list(base)[1:]:base.remove(s)
    missing=set(c for d in state['weeks'].values() for c in d.get('missing_late',[]))
    semester='2025-2026学年第一学期'
    complete=set(colleges)
    for d in state['weeks'].values():complete &= set(d['received']['attendance']) & set(d['received']['late'])
    base.active.title=f'{semester}{scope}学生课堂纪律登记表';order=fill(base.active,ref,groups,submitted,colleges,f'{semester}{scope}学生课堂纪律登记表',True,missing,complete)
    tw=load_workbook(weekly);template=tw.active
    for n in weeks:
        chinese={1:'一',2:'二',3:'三',4:'四',5:'五',6:'六',7:'七',8:'八',9:'九',10:'十'}
        if n not in chinese:raise ValueError('当前周次命名支持第一至第十周')
        s=base.create_sheet(f'第{chinese[n]}周')
        s.sheet_format=copy.copy(template.sheet_format)
        for row in template:
            for c in row:
                if c.row<=2:copy_style(c,s.cell(c.row,c.column,c.value))
        for k,d in template.column_dimensions.items():s.column_dimensions[k]=copy.copy(d)
        for r in (1,2):s.row_dimensions[r]=copy.copy(template.row_dimensions[r])
        d=state['weeks'][str(n)];g,sub=aggregate([d]);complete_week=set(d['received']['attendance']) & set(d['received']['late'])
        fill(s,template,g,sub,colleges,f'{semester}第{n}周学生课堂纪律登记表',False,d.get('missing_late',[]),complete_week)
    file=output/f'{semester}{scope}学生课堂纪律登记表.xlsx';tmp=output/'写入中.xlsx';base.save(tmp)
    verified=load_workbook(tmp,data_only=True)
    for s in verified:
        for row in s:
            for c in row:
                if c.value=='未及时报送' and c.column==2:
                    assert c.font.color.rgb=='00FF0000' or c.font.color.rgb=='FFFF0000'
                    assert any(c.coordinate in m and m.min_col==2 and m.max_col==(6 if s==verified.worksheets[0] else 5) for m in s.merged_cells.ranges)
    total=sum(v['hours'] for v in groups.values());actual=sum(c.value for (c,) in verified.active.iter_rows(min_row=3,min_col=4,max_col=4) if isinstance(c.value,(int,float)))
    assert total==actual
    verified.close();os.replace(tmp,file)
    st=output/'台账写入中.json';st.write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8');os.replace(st,statepath)
    report={'output':str(file),'imported_week':week,'weeks':weeks,'total_hours':total,'abnormal_students':len(groups),'college_order':order,'filename_fallback':data['filename_fallback'],'evidence_count':len(data['events']),'source_files':data['files']}
    (output/f'第{week}周运行报告.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--root',required=True);p.add_argument('--output',required=True);p.add_argument('--week',type=int,required=True);a=p.parse_args()
    print(json.dumps(run(a.root,a.output,a.week),ensure_ascii=False))
