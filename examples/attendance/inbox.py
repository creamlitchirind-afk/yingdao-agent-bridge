"""Transactional weekly inbox; invoked by a Yingdao code flow."""
from pathlib import Path
import argparse, contextlib, hashlib, json, os, re, shutil, sys, uuid
from collections import Counter
from datetime import datetime
sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine
from openpyxl import load_workbook

NAMES = {i: '第'+v+'周' for i,v in enumerate(['一','二','三','四','五','六','七','八','九','十'],1)}

def write_json(p, data):
    tmp=p.with_suffix('.tmp')
    tmp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    os.replace(tmp,p)

def read_json(p): return json.loads(p.read_text(encoding='utf-8'))

def hashes(folder):
    return {str(p.relative_to(folder)):hashlib.sha256(p.read_bytes()).hexdigest()
            for p in folder.rglob('*') if p.is_file() and not p.name.startswith('~$')}

@contextlib.contextmanager
def lock(base):
    import msvcrt
    with (base/'.运行锁').open('a+b') as f:
        f.seek(0); f.write(b'0'); f.flush(); f.seek(0)
        try: msvcrt.locking(f.fileno(),msvcrt.LK_NBLCK,1)
        except OSError: raise RuntimeError('已有导入任务正在运行，请等待完成。')
        try: yield
        finally:
            f.seek(0); msvcrt.locking(f.fileno(),msvcrt.LK_UNLCK,1)

def verify(output, ledger):
    """Read output back; compare every student, order, status, merge and note."""
    w=load_workbook(output,data_only=True)
    weeks=sorted(map(int,ledger['weeks']))
    assert w.sheetnames[1:]==[NAMES[n] for n in weeks]
    checks=[]
    for index,s in enumerate(w):
        ds=list(ledger['weeks'].values()) if index==0 else [ledger['weeks'][str(weeks[index-1])]]
        expected={}; complete=None; missing=set()
        for d in ds:
            both=set(d['received']['attendance'])&set(d['received']['late'])
            complete=both if complete is None else complete&both
            missing.update(d['missing_late'])
            for e in d['events']:
                k=(e['college'],e['class'],e['name'])
                expected[k]=expected.get(k,0)+e['hours']
        actual={}; college=None; categories=[]; group=[]; last=6 if index==0 else 5
        for r in range(3,s.max_row):
            if s.cell(r,1).value:
                if group: assert group==sorted(group)
                group=[];college=s.cell(r,1).value
                count=sum(k[0]==college for k in expected)
                categories.append((1,count) if count else (0,0) if college in complete else (2,0))
            klass,name,h=s.cell(r,2).value,s.cell(r,3).value,s.cell(r,4).value
            note=s.cell(r,last)
            if name:
                k=(college,klass,name);assert k not in actual
                actual[k]=h;group.append((-h,name,klass))
                if college in missing:assert '未报送，该名单为学工处考勤结果' in str(note.value)
            else:
                assert klass==('全勤' if college in complete else '未及时报送')
                end='D' if klass=='全勤' else chr(64+last)
                assert f'B{r}:{end}{r}' in set(map(str,s.merged_cells.ranges))
                if klass!='全勤':assert s.cell(r,2).font.color.rgb[-6:]=='FF0000'
            if note.value:assert note.font.color.rgb[-6:]=='FF0000'
        if group:assert group==sorted(group)
        assert actual==expected and categories==sorted(categories)
        checks.append({'sheet':s.title,'students':len(actual),'hours':sum(actual.values())})
    w.close()
    return checks

def archive_committed(base, current, update_preview=True):
    """Complete the rename after a crash between publication and marking."""
    if not current:return
    version=base/'总表版本'/current['version']
    manifest=read_json(version/'导入记录.json')
    source=base/'本周周表'/manifest['source_name']
    target=source.with_name(source.name+'-已导入')
    latest=base/'当前总表';latest.mkdir(exist_ok=True)
    staged=latest/'最新总表.tmp.xlsx'
    try:
        if not update_preview:raise InterruptedError()
        shutil.copy2(current['output'],staged)
        os.replace(staged,latest/'最新总表.xlsx')
        alias_status={'updated':True,'output':current['output']}
    except InterruptedError:
        alias_status=None
    except PermissionError:
        # The immutable, verified version is already committed. An open preview
        # must not block import or cause the same week to be counted twice.
        alias_status={'updated':False,'reason':'最新总表被占用，打开的旧标签页不会刷新，请查看本次版本文件。','output':current['output']}
    if alias_status is not None:write_json(base/'最新总表更新状态.json',alias_status)
    if source.exists():
        if hashes(source)!=manifest['source_hashes']:raise RuntimeError('已提交周目录内容发生变化，请保留材料并人工核对：'+str(source))
        if target.exists():raise RuntimeError('已导入目录同名冲突：'+str(target))
        try:source.rename(target)
        except PermissionError:
            # Open input workbooks can prevent renaming their parent directory.
            # The committed manifest is the import marker until the lock clears.
            return False
    elif not target.exists():raise RuntimeError('已提交材料目录与归档目录均缺失：'+manifest['source_name'])
    if hashes(target)!=manifest['source_hashes']:raise RuntimeError('已导入材料已被修改：'+str(target))
    return True

def verify_graphical_outputs(base, source, weekdata):
    """Gate publication on the native flows' independently computed student hours.

    Only enabled by the integrated graphical main flow. Direct code entry and
    historical callers retain their existing behavior.
    """
    reports=[]
    for filename,kind in [('图形课堂事件.json','课堂'),('图形晚点名事件.json','晚点名')]:
        rows=json.loads((base/filename).read_text(encoding='utf-8-sig'))
        if not isinstance(rows,list):raise ValueError('图形事件必须为列表：'+filename)
        actual=Counter();expected=Counter();count=0
        for row in rows:
            if not isinstance(row,list) or len(row)!=5:raise ValueError('图形事件结构错误：'+filename)
            college,klass,name,hours,path=row
            if not isinstance(path,str):raise ValueError('图形事件缺少材料路径')
            material=Path(path).resolve()
            if material.parent.parent!=source.resolve():continue
            if not isinstance(hours,(int,float)) or isinstance(hours,bool) or hours<=0:
                raise ValueError('图形事件学时无效：'+filename)
            key=(engine.normalized(college),engine.normalized(klass),engine.normalized(name),material.name)
            actual[key]+=hours;count+=1
        for event in weekdata['events']:
            if event['kind']==kind:
                key=(event['college'],event['class'],event['name'],Path(event['file']).name)
                expected[key]+=event['hours']
        if actual!=expected:
            differing=[k for k in actual.keys()|expected.keys() if actual[k]!=expected[k]]
            raise ValueError(f'第{weekdata["week"]}周{kind}图形计算与快照核验不一致，未提交；差异：'+str(differing[:5]))
        reports.append({'kind':kind,'records':count,'student_source_keys':len(actual),'hours':sum(actual.values()),'matched':True})
    return reports

def run(base, progress=None, graphical_evidence=False):
    notify=progress or (lambda message: None)
    base=Path(base).resolve();base.mkdir(parents=True,exist_ok=True)
    inbox=base/'本周周表';inbox.mkdir(exist_ok=True)
    versions=base/'总表版本';versions.mkdir(exist_ok=True)
    pointer=base/'当前总表.json'
    with lock(base):
        notify('[开始] 检查历史版本与中断恢复记录')
        current=read_json(pointer) if pointer.exists() else None
        archive_committed(base,current)
        committed_weeks=set();pending_marks=[]
        cursor=current
        while cursor:
            manifest=read_json(versions/cursor['version']/'导入记录.json')
            committed_weeks.add(manifest['week'])
            if archive_committed(base,cursor,False) is False:pending_marks.append(manifest['source_name'])
            previous=manifest.get('previous_version')
            cursor={'version':previous,'output':current['output']} if previous else None
        pending={}
        for p in inbox.iterdir():
            if not p.is_dir() or p.name.endswith('-已导入'):continue
            week=next((n for n,name in NAMES.items() if p.name in (name,f'第{n}周')),None)
            if week is None:raise ValueError('无法识别周文件夹：'+p.name+'；请命名为第一周至第十周')
            if week in committed_weeks:continue
            if week in pending:raise ValueError('同一周存在两个待导入文件夹')
            pending[week]=p
        ledger=read_json(base/'总表版本'/current['version']/'结果/汇总台账.json') if current else {'weeks':{}}
        expected=max(map(int,ledger['weeks']),default=0)+1
        if sorted(pending)!=list(range(expected,expected+len(pending))):
            raise ValueError(f'需要从第{expected}周连续导入，当前待导入周为{sorted(pending)}；已有周不重复导入。')
        done=[]
        notify('[完成] 待导入周识别：'+str(sorted(pending)))
        for week,source in sorted(pending.items()):
            notify(f'[开始] 第{week}周：检查材料并创建独立版本')
            for folder in ('考勤表','晚点名'):
                if not (source/folder).is_dir():raise ValueError(str(source/folder)+'不存在；未收到材料时保留空文件夹。')
            source_hashes=hashes(source)
            if not any(p.suffix.lower()=='.xlsx' and not p.name.startswith('~$') for p in source.rglob('*')):
                raise ValueError('本周没有Excel材料：'+str(source))
            if source.with_name(source.name+'-已导入').exists():raise ValueError('已导入目录已存在：'+source.name)
            version=NAMES[week]+'_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:6]
            dest=versions/version;dest.mkdir()
            inputs=dest/'材料快照';inputs.mkdir()
            shutil.copytree(base/'基础模板/01_起始总表',inputs/'01_起始总表')
            shutil.copy2(base/'基础模板/课时分组.json',inputs/'课时分组.json')
            shutil.copytree(source,inputs/'02_已提交材料'/f'第{week}周',ignore=shutil.ignore_patterns('~$*'))
            output=dest/'结果';output.mkdir()
            write_json(output/'汇总台账.json',ledger)
            try:
                notify(f'[完成] 第{week}周材料快照；开始汇总')
                result=engine.run(inputs,output,week,progress=notify)
                notify(f'[完成] 第{week}周汇总；开始逐人及格式核验')
                newledger=read_json(output/'汇总台账.json')
                checks=verify(result['output'],newledger)
                graphical_checks=verify_graphical_outputs(base,source,newledger['weeks'][str(week)]) if graphical_evidence else None
                if graphical_checks:notify(f'[节点完成] 第{week}周课堂与晚点名图形结果已逐人核对，允许提交')
                if hashes(source)!=source_hashes:raise RuntimeError('运行期间输入发生变化，未提交此次结果。')
                manifest={'week':week,'previous_version':current['version'] if current else None,'source_name':source.name,'source_hashes':source_hashes,'result':result,'checks':checks,'status':'verified'}
                if graphical_checks:manifest['graphical_checks']=graphical_checks
                write_json(dest/'导入记录.json',manifest)
                current={'version':version,'week':week,'output':result['output'],'ledger':str(output/'汇总台账.json')}
                write_json(pointer,current)
                weekdata=newledger['weeks'][str(week)]
                for college in result['college_order']:
                    entries=[e for e in weekdata['events'] if e['college']==college]
                    people=len({(e['class'],e['name']) for e in entries})
                    complete=all(college in weekdata['received'][kind] for kind in ('attendance','late'))
                    state='异常名单已写入' if entries else '全勤' if complete else '未及时报送'
                    extra='；缺晚点名，已加未报送备注' if entries and college in weekdata['missing_late'] else ''
                    notify(f'[学院导入完成] {college}第{week}周：{state}，{people}人、{sum(e["hours"] for e in entries)}学时{extra}')
                notify(f'[完成] 第{week}周核验与版本提交：累计{result["abnormal_students"]}人、{result["total_hours"]}学时')
            except Exception as exc:
                notify(f'[错误] 第{week}周处理失败：{exc}')
                write_json(dest/'失败记录.json',{'error':str(exc),'published':False})
                raise
            if archive_committed(base,current) is False:pending_marks.append(source.name)
            ledger=newledger;done.append({'week':week,'version':version,'checks':checks})
            notify(f'[完成] 第{week}周归档处理（文件占用时记录待改名）')
        result={'status':'completed' if done else 'no_pending','imported':done,'current':current,
                'latest_preview':read_json(base/'最新总表更新状态.json') if current else None,
                'pending_folder_renames':pending_marks}
        write_json(base/'最近运行结果.json',result)
        warnings=[]
        if result['latest_preview'] and not result['latest_preview']['updated']:warnings.append(result['latest_preview']['reason'])
        if pending_marks:warnings.append('待补改名：'+', '.join(pending_marks))
        for d in ledger['weeks'].values():
            if d.get('missing_late'):warnings.append('第'+str(d['week'])+'周缺晚点名：'+', '.join(d['missing_late']))
        result['warnings']=warnings
        notify('[业务问题汇总] '+('；'.join(warnings) if warnings else '无异常、无待处理警告'))
        notify('[业务总日志] 本次新导入'+str(len(done))+'周；历史共'+str(len(ledger['weeks']))+'周；问题/提醒'+str(len(warnings))+'项；当前结果：'+str(current['output'] if current else '无'))
        write_json(base/'最近运行结果.json',result)
        return result

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--base',required=True);a=p.parse_args()
    print(json.dumps(run(a.base),ensure_ascii=False))
