import importlib.util,json,tempfile,unittest,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'examples'))
import shop_orders as d
from shop_orders_plan import make_plan,remap_variables,HEADERS

class ShopTests(unittest.TestCase):
 def fixture(self,base):
  raw=[[
   ['O1','短袖T恤','¥9.99','2026-01-01','完成'],
   ['O2','短袖T恤','¥20','2026-01-01','取消']],
   [['O3','短袖T恤','￥188','2026-01-02','完成'],
    ['O1','短袖T恤','¥9.99','2026-01-01','完成']]]
  r={'raw_output':str(base/'raw.json'),'output':str(base/'out.xlsx'),'pages':2,'product':'短袖T恤'}
  Path(r['raw_output']).write_text(json.dumps(raw,ensure_ascii=False),encoding='utf-8')
  from openpyxl import Workbook
  w=Workbook();s=w.active;s.append(HEADERS);s.append(['O1','短袖T恤',9.99,'2026-01-01','完成']);s.append(['O3','短袖T恤',188,'2026-01-02','完成']);w.save(r['output']);w.close()
  return r,raw
 def test_filter_order_dedup_and_money(self):
  with tempfile.TemporaryDirectory() as tmp:
   r,_=self.fixture(Path(tmp));result=d.verify(r)
   self.assertEqual((result['source_rows'],result['record_count'],result['total_amount']),(4,2,197.99))
 def test_wrong_product_blocked(self):
  with tempfile.TemporaryDirectory() as tmp:
   r,raw=self.fixture(Path(tmp));raw[0][0][1]='其他商品';Path(r['raw_output']).write_text(json.dumps(raw),encoding='utf-8')
   with self.assertRaisesRegex(ValueError,'product'):d.verify(r)
 def test_repeated_page_blocked(self):
  with tempfile.TemporaryDirectory() as tmp:
   r,raw=self.fixture(Path(tmp));raw[1]=raw[0];Path(r['raw_output']).write_text(json.dumps(raw),encoding='utf-8')
   with self.assertRaisesRegex(ValueError,'repeated'):d.verify(r)
 def test_wrong_excel_order_blocked(self):
  with tempfile.TemporaryDirectory() as tmp:
   r,_=self.fixture(Path(tmp));from openpyxl import load_workbook
   w=load_workbook(r['output']);s=w.active;a=list(s.values)[1:];s.delete_rows(2,2)
   for row in reversed(a):s.append(row)
   w.save(r['output']);w.close()
   with self.assertRaisesRegex(ValueError,'differs'):d.verify(r)
 def test_plan_is_native_and_parameterized(self):
  a=make_plan('短袖T恤',2,Path('a.xlsx'),Path('a.json'));b=make_plan('篮球鞋',1,Path('b.xlsx'),Path('b.json'))
  banned={'programing.snippet','process.run','web.browser.execute_javascript'}
  self.assertFalse(banned&{x['prototype'] for x in a['blocks']})
  self.assertNotIn('expression',json.dumps(a))
  self.assertIn('篮球鞋',json.dumps(b,ensure_ascii=False));self.assertIn('b.xlsx',json.dumps(b))
 def test_session_parameters_cannot_change(self):
  with tempfile.TemporaryDirectory() as tmp:
   d.prepare('短袖T恤',2,Path(tmp));self.assertEqual(d.prepare('短袖T恤',2,Path(tmp))['status'],'prepared')
   with self.assertRaisesRegex(ValueError,'differ'):d.prepare('篮球鞋',1,Path(tmp))
 def test_failed_worker_cannot_be_relaunched(self):
  with tempfile.TemporaryDirectory() as tmp:
   base=Path(tmp);d.prepare('短袖T恤',2,base)
   (base/'worker.json').write_text(json.dumps({'state':'failed','error':'connection failed'}),encoding='utf-8')
   result=d.background('launch','短袖T恤',2,base)
   self.assertEqual(result['status'],'blocked')
   self.assertEqual(json.loads((base/'worker.json').read_text())['state'],'failed')
 def test_verified_session_is_not_run_again(self):
  with tempfile.TemporaryDirectory() as tmp:
   base=Path(tmp);r=d.prepare('短袖T恤',2,base);r['status']='verified'
   (base/'session.json').write_text(json.dumps(r),encoding='utf-8')
   self.assertEqual(d.background('launch','短袖T恤',2,base)['status'],'verified')
 def test_all_pages_requires_terminal_evidence(self):
  with tempfile.TemporaryDirectory() as tmp:
   r,_=self.fixture(Path(tmp));r['pages']=0
   with self.assertRaisesRegex(ValueError,'last-page'):d.verify(r)
   proof=Path(r['raw_output']).with_name('raw_分页核验.json')
   proof.write_text(json.dumps([False,'ant-pagination-next',2]),encoding='utf-8')
   with self.assertRaisesRegex(ValueError,'last-page'):d.verify(r)
   proof.write_text(json.dumps([True,'ant-pagination-next ant-pagination-disabled',2]),encoding='utf-8')
   self.assertEqual(d.verify(r)['page_count'],2)
 def test_workbook_is_visible_before_write_and_stays_open(self):
  blocks=make_plan('短袖T恤',0,Path('a.xlsx'),Path('a.json'))['blocks']
  launch=next(i for i,b in enumerate(blocks) if b['prototype']=='system.run_or_open')
  writing=next(i for i,b in enumerate(blocks) if b['prototype']=='win32.send_keys' and b['inputs']['keys']=='^v')
  maximize=next(i for i,b in enumerate(blocks) if b['prototype']=='win32.window.set_state' and b['inputs'].get('window')=={'variable':'report_window'})
  self.assertLess(launch,maximize);self.assertLess(maximize,writing)
  self.assertTrue(blocks[launch]['inputs']['application_path'].endswith('.xlsx'))
  save=next(i for i,b in enumerate(blocks) if b['prototype']=='win32.send_keys' and b['inputs']['keys']=='^s')
  self.assertLess(writing,save)
  self.assertFalse({'excel.close','win32.window.close'}&{b['prototype'] for b in blocks})
 def test_wps_native_datetime_compares_exact_timestamp(self):
  from datetime import datetime
  from openpyxl import load_workbook
  with tempfile.TemporaryDirectory() as tmp:
   r,raw=self.fixture(Path(tmp));raw[0][0][3]='2026-01-01 00:00:00';raw[1][0][3]='2026-01-02 00:00:00'
   Path(r['raw_output']).write_text(json.dumps(raw,ensure_ascii=False),encoding='utf-8')
   w=load_workbook(r['output']);w.active['D2']=datetime(2026,1,1);w.active['D3']=datetime(2026,1,2);w.save(r['output']);w.close()
   self.assertEqual(d.verify(r)['record_count'],2)
   w=load_workbook(r['output']);w.active['D2']=datetime(2026,1,1,0,1);w.save(r['output']);w.close()
   with self.assertRaisesRegex(ValueError,'differs'):d.verify(r)
 def test_runtime_variable_mapping_preserves_business_text(self):
  plan={'flow_id':'demo','blocks':[{'prototype':'example','inputs':{'browser':{'variable':'shop_browser'},'text':'shop_browser'},'output':'shop_browser'}]}
  mapped=remap_variables(plan,{'shop_browser':'web_page'})
  self.assertEqual(mapped['blocks'][0]['inputs']['browser'],{'variable':'web_page'})
  self.assertEqual(mapped['blocks'][0]['inputs']['text'],'shop_browser')
  self.assertEqual(mapped['blocks'][0]['output'],'web_page')
 def test_new_sessions_do_not_reuse_open_workbook_name(self):
  with tempfile.TemporaryDirectory() as tmp:
   first=d.prepare('短袖T恤',0,Path(tmp)/'one');second=d.prepare('短袖T恤',0,Path(tmp)/'two')
   self.assertNotEqual(Path(first['output']).name,Path(second['output']).name)
   self.assertEqual(d.prepare('短袖T恤',0,Path(tmp)/'one')['output'],first['output'])
 def test_layout_asset_is_blank_and_does_not_shift_home_to_row_two(self):
  from openpyxl import load_workbook
  w=load_workbook(ROOT/'examples/assets/shop-report-empty.xlsx')
  try:
   self.assertIsNone(w.active.freeze_panes)
   self.assertFalse(any(c.value is not None for row in w.active for c in row))
  finally:w.close()
 def test_product_text_does_not_become_a_file_path(self):
  with tempfile.TemporaryDirectory() as tmp:
   session=Path(tmp)/'session';product='../款式/测试\\名称:*?'
   r=d.prepare(product,0,session)
   self.assertEqual(r['product'],product)
   self.assertEqual(Path(r['output']).parent,session.resolve())
   self.assertFalse(set('<>:"/\\|?*')&set(Path(r['output']).name))

if __name__=='__main__':unittest.main()
