import importlib.util, json, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('demo',ROOT/'examples/attendance_video.py')
d=importlib.util.module_from_spec(spec);spec.loader.exec_module(d)

class DemoTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        self.fixture_dir=tempfile.TemporaryDirectory()
        self.addCleanup(self.fixture_dir.cleanup)
        self.source=Path(self.fixture_dir.name)/'input.xlsx'
        book=Workbook();sheet=book.active;sheet.append(d.HEADERS)
        for row in [['学院甲','一班','模拟甲',0],['学院乙','二班','模拟乙',1],['学院甲','一班','模拟丙',3]]:sheet.append(row)
        book.save(self.source);book.close()

    def test_active_worker_blocks_duplicate_launch(self):
        with tempfile.TemporaryDirectory() as temp:
            session=Path(temp)
            with patch.object(d,'lock_file',return_value=None),patch.object(d.subprocess,'Popen') as spawn:
                result=d.background('launch',self.source,session)
                self.assertEqual(result['status'],'busy')
                spawn.assert_not_called()

    def test_completed_worker_queries_existing_run_only(self):
        with tempfile.TemporaryDirectory() as temp:
            session=Path(temp)
            with patch.object(d,'lock_file') as lock,patch.object(d,'execute') as execute:
                execute.side_effect=[{'run_id':'existing-run'}, {'status':'verified'}]
                d.background('worker',self.source,session)
                self.assertEqual([call.args[0] for call in execute.call_args_list],['prepare','check'])
                lock.return_value.close.assert_called_once()

    def test_native_recipe_contract_and_readback(self):
        rows=d.read_rows(self.source)
        plan=d.make_plan(self.source,Path('demo-output.xlsx'),rows)
        def check(value):
            if isinstance(value,dict):
                if 'expression' in value:
                    self.assertIn(value['expression'],['True','False','0','1','2','3',str(len(rows))])
                for item in value.values():check(item)
            elif isinstance(value,list):
                for item in value:check(item)
        check(plan)
        actual={'blocks':[{'prototype_name':x['prototype'],'fields':{k:{'value':v} for k,v in x['inputs'].items()},'outputs':x['outputs']} for x in plan['blocks']]}
        self.assertEqual(d.compare_plan(plan,actual),[])
        actual['blocks'][1]['fields']['read_way']['value']='used_range'
        self.assertTrue(d.compare_plan(plan,actual))

    def test_wrong_app_never_creates(self):
        with tempfile.TemporaryDirectory() as temp:
            session=Path(temp)
            d.execute('prepare',self.source,session)
            with patch.object(d,'configure'),patch.object(d.bridge,'cli_current',return_value={'data':{'app':{'app_id':'WRONG'}}}),patch.object(d.bridge,'create_flow') as create:
                with self.assertRaises(RuntimeError): d.execute('build',self.source,session)
                create.assert_not_called()

    def test_output_exactness(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as temp:
            session=Path(temp);record=d.execute('prepare',self.source,session)
            book=Workbook();sheet=book.active;sheet.append(d.HEADERS)
            for row in record['expected']:sheet.append(row)
            book.save(record['output']);book.close()
            self.assertEqual(d.verify_output(record)['total_hours'],4)
            book=Workbook();sheet=book.active;sheet.append(d.HEADERS)
            for row in reversed(record['expected']):sheet.append(row)
            book.save(record['output']);book.close()
            with self.assertRaises(ValueError):d.verify_output(record)

    def test_empty_filtered_result_is_valid(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'all-present.xlsx';book=Workbook();s=book.active;s.append(d.HEADERS);s.append(['学院甲','一班','模拟甲',0]);book.save(source);book.close()
            record=d.execute('prepare',source,Path(temp)/'session')
            book=Workbook();book.active.append(d.HEADERS);book.save(record['output']);book.close()
            self.assertEqual(d.verify_output(record)['record_count'],0)

if __name__=='__main__':unittest.main()
