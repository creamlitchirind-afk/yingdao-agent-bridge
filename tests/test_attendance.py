from pathlib import Path
import tempfile,sys,unittest
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'examples/attendance'))
from engine import run,lesson_events
from types import SimpleNamespace

class AttendanceTests(unittest.TestCase):
 def test_lesson_boundary(self):
  class Sheet:
   def cell(self,r,c):return SimpleNamespace(value={3:'○',5:'#'}.get(c),coordinate=f'{c}:{r}')
  self.assertEqual(sum(h for h,_ in lesson_events(Sheet(),6,[3,4,5],[[3,4,5]])),.5)
  self.assertEqual(sum(h for h,_ in lesson_events(Sheet(),6,[3,4,5],[[3,4],[5]])),1)

 def test_three_weeks_and_repeat(self):
  data=ROOT/'examples/attendance/data'
  with tempfile.TemporaryDirectory() as tmp:
   for week,expected in [(1,(80,162)),(2,None),(3,(124,483.5))]:
    result=run(data,Path(tmp),week)
    if week in (1,3):self.assertEqual((result['abnormal_students'],result['total_hours']),expected)
   path=Path(result['output']);w=load_workbook(path)
   self.assertEqual(w.sheetnames[1:],['第一周','第二周','第三周'])
   for sheet,count,hours in [('第一周',80,162),('第二周',94,176.5),('第三周',89,145)]:
    s=w[sheet];values=[s.cell(r,4).value for r in range(3,s.max_row) if isinstance(s.cell(r,4).value,(int,float))]
    self.assertEqual((len(values),sum(values)),(count,hours))
   s=w['第三周'];last=s.max_row-1
   self.assertEqual(s.cell(last,2).value,'未及时报送')
   self.assertIn(f'B{last}:E{last}',{str(m) for m in s.merged_cells.ranges})
   self.assertEqual(s.cell(last,2).font.color.rgb[-6:],'FF0000')
   before={s.title:[[c.value for c in row] for row in s] for s in w};w.close()
   result=run(data,Path(tmp),3);after=load_workbook(result['output'])
   self.assertEqual(before,{s.title:[[c.value for c in row] for row in s] for s in after});after.close()

if __name__=='__main__':unittest.main()
